"""
Generate the one-page abstract (.docx) for submission.

It reads the REAL numbers out of results/benchmark.json, so the abstract can
never drift from what was actually measured -- re-run it after any change to
the experiments and the document updates itself.

    python scripts/make_abstract.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document                                      # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                  # noqa: E402
from docx.oxml.ns import qn                                    # noqa: E402
from docx.shared import Pt, RGBColor                           # noqa: E402

from src import config                                         # noqa: E402

TITLE = "Cross-Encoder Reranking Pipeline for Patent Prior-Art Search"
BYLINE = "Roll Numbers: 22PD32, 22PD12   |   Team members: Shivnaran S, Diravina E I"
COURSE = "Information Retrieval & Web Search Laboratory"

ACCENT = RGBColor(0x2A, 0x33, 0x6B)
INK = RGBColor(0x1A, 0x1A, 0x1A)


# --------------------------------------------------------------------------- #
def load_numbers() -> dict:
    stats = json.loads((config.DATA_DIR / "dataset_stats.json").read_text(encoding="utf-8"))
    bench_path = config.RESULTS_DIR / "benchmark.json"
    bench = json.loads(bench_path.read_text(encoding="utf-8")) if bench_path.exists() else None
    return {"stats": stats, "bench": bench}


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = normal.paragraph_format
    pf.space_after = Pt(1.5)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0

    # The brief says STRICTLY one page, so the margins are tight and
    # scripts/make_abstract.py is the only place the text lives -- keep it short.
    for section in doc.sections:
        section.top_margin = Pt(32)
        section.bottom_margin = Pt(28)
        section.left_margin = Pt(42)
        section.right_margin = Pt(42)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = ACCENT


def body(doc: Document, text: str, bullet: bool = False) -> None:
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    if bullet:
        p.paragraph_format.left_indent = Pt(14)
    run = p.add_run(text)
    run.font.size = Pt(9.5)


def build(out_path: Path) -> None:
    data = load_numbers()
    stats, bench = data["stats"], data["bench"]
    doc = Document()
    style_document(doc)

    # ---- title ------------------------------------------------------------ #
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(COURSE)
    r.font.size = Pt(9)
    r.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(BYLINE)
    r.font.size = Pt(9)
    r.bold = True

    # ---- problem statement ------------------------------------------------ #
    heading(doc, "Problem Statement")
    body(doc,
         "Before a patent is granted, an examiner must find prior art: earlier documents "
         "that already disclose the claimed invention. Missing it leads to invalid patents "
         "and costly litigation. Three things make the retrieval hard. Queries are patent "
         f"claims, not keywords (average {stats['mean_query_words']:.0f} words of technical "
         "and legal prose); applicants coin their own terminology, so two documents "
         "describing the same mechanism may share almost no vocabulary, which defeats "
         "lexical matching; and the most accurate model, a cross-encoder, must read the "
         "query and a document together, so it cannot be applied to a whole collection. "
         "Given a claim as the query, rank earlier documents so that the prior art an "
         "examiner would cite appears as high as possible.")

    # ---- objective -------------------------------------------------------- #
    heading(doc, "Objective")
    body(doc,
         "Build and evaluate a two-stage retrieve-and-rerank pipeline: a fast high-recall "
         "first stage narrows the collection to a shortlist, and a high-precision "
         "transformer reranks only that shortlist. Compare lexical, dense and hybrid "
         "first stages fairly; fine-tune a cross-encoder using examiner citations and "
         "mined hard negatives; and quantify through ablations and error analysis what "
         "each component contributes and what it costs in latency.")

    # ---- dataset ---------------------------------------------------------- #
    heading(doc, "Dataset Used")
    body(doc,
         "PatentMatch, derived from European Patent Office examiner search reports. When "
         "an examiner rejects a claim for lack of novelty, the report cites the exact "
         "paragraph that disclosed it; those citations are used directly as relevance "
         f"judgements, so the ground truth is expert human judgement. The collection has "
         f"{stats['n_documents']:,} prior-art passages, {stats['n_queries']:,} claim queries "
         f"and {stats['n_qrels']:,} judgements, plus {stats['n_hard_negatives']} passages "
         "examiners inspected and rejected, which serve as gold-standard hard negatives. "
         "Because near-duplicate claims of one application cite the same prior art, a "
         "random split would leak; we assign whole connected components of the "
         f"query-document graph to one split ({stats['queries_per_split']['train']:,} train / "
         f"{stats['queries_per_split']['dev']:,} dev / "
         f"{stats['queries_per_split']['test']:,} test), so no test query shares a relevant "
         "document with any training query. HUPD was dropped: it has no citation field and "
         "the PatentsView API is no longer openly reachable, so labels cannot be derived "
         "from it.")

    # ---- methods ---------------------------------------------------------- #
    heading(doc, "Methods / Algorithms Used")
    body(doc,
         "Preprocessing: tokenisation, English and patent-specific stopword removal "
         "(wherein, comprising, said), Porter stemming, applied identically to queries "
         "and documents.", bullet=True)
    body(doc,
         "Stage 1 (recall): TF-IDF with SMART ltc weighting and cosine; Okapi BM25 "
         "implemented from scratch with per-term score attribution; a dense bi-encoder "
         "(all-MiniLM-L6-v2, mean pooling) in a FAISS IndexFlatIP for exact cosine "
         "search; and a hybrid fusing BM25 and dense rankings by Reciprocal Rank Fusion, "
         "which uses ranks only and so needs no score calibration.", bullet=True)
    body(doc,
         "Stage 2 (precision): a pairwise Learning-to-Rank model over nine ranking "
         "features with a RankNet loss written in numpy, as a cheap control; and a "
         "cross-encoder (ms-marco-MiniLM-L-6-v2) zero-shot and then fine-tuned on "
         "PatentMatch. Long passages are truncated or scored by MaxP over overlapping "
         "windows.", bullet=True)
    body(doc,
         "Hard negative mining: the real first stage is run on each training query and "
         "its top-ranked non-cited documents become negatives, sampled across the pool, "
         "so the reranker trains on the decisions it faces at inference.", bullet=True)
    body(doc,
         "Evaluation: Recall@k, Precision@k, MRR, MAP, nDCG@k from scratch, with a paired "
         "randomisation test and per-stage latency. Demonstrated through a seven-page "
         "Streamlit application and verified by 41 tests against hand-computed values.",
         bullet=True)

    # ---- results ---------------------------------------------------------- #
    heading(doc, "Results (Metrics Used)")
    if bench:
        methods = bench["methods"]
        cut = bench["primary_cutoff"]
        order = [k for k in ["tfidf", "bm25", "dense", "hybrid", "hybrid_ltr",
                             "hybrid_ce", "hybrid_ce_ft"] if k in methods]
        best_key = max(order, key=lambda k: methods[k]["metrics"][f"ndcg@{cut}"])
        best, base = methods[best_key], methods.get("bm25")
        bm = best["metrics"]

        # The brief's format asks for an explicit "[Metric]: [final score]" line,
        # so the headline numbers are stated plainly before the full table.
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(
            f"Final scores ({best['label']}):   "
            f"nDCG@{cut} = {bm[f'ndcg@{cut}']:.4f}   |   "
            f"MRR = {bm['mrr']:.4f}   |   "
            f"MAP = {bm['map']:.4f}   |   "
            f"Recall@100 = {bm['recall@100']:.4f}"
        )
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = ACCENT

        table = doc.add_table(rows=1, cols=7)
        table.style = "Light Grid Accent 1"
        table.autofit = False
        # Fixed widths: the method names need the room, the numbers do not.
        # Without this a long label wraps onto a second line and the table
        # grows past the one-page limit.
        widths = [Pt(132), Pt(52), Pt(44), Pt(44), Pt(56), Pt(60), Pt(44)]
        hdr = ["Method", f"nDCG@{cut}", "MRR", "MAP", "R@10", "R@100", "ms/q"]

        def fill(cells, values, bold: bool) -> None:
            for cell, text, width in zip(cells, values, widths):
                cell.width = width
                cell.text = ""
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.space_before = Pt(0)
                run = para.add_run(text)
                run.font.size = Pt(8)
                run.bold = bold

        fill(table.rows[0].cells, hdr, True)
        for key in order:
            m = methods[key]
            met = m["metrics"]
            fill(table.add_row().cells,
                 [m["label"], f"{met[f'ndcg@{cut}']:.4f}", f"{met['mrr']:.4f}",
                  f"{met['map']:.4f}", f"{met['recall@10']:.4f}",
                  f"{met['recall@100']:.4f}", f"{m['mean_latency_ms']:.0f}"],
                 key == best_key)

        doc.add_paragraph().paragraph_format.space_after = Pt(1)

        gain = best["metrics"][f"ndcg@{cut}"] - base["metrics"][f"ndcg@{cut}"]
        rel = gain / base["metrics"][f"ndcg@{cut}"] * 100 if base["metrics"][f"ndcg@{cut}"] else 0
        def sig_for(method_key: str, metric: str) -> dict | None:
            return next((r for r in bench.get("significance", [])
                         if r["method"] == method_key and r["metric"] == metric), None)

        sig = sig_for(best_key, f"ndcg@{cut}")
        sig_text = ""
        if sig:
            sig_text = (f" A paired randomisation test puts this at p = {sig['p_value']:.4f}, "
                        f"so it is statistically significant."
                        if sig["significant_at_05"] else
                        f" A paired randomisation test gives p = {sig['p_value']:.4f}, so on "
                        f"this query sample the nDCG gain is NOT statistically significant; "
                        f"we report it as measured rather than overstating it.")

        hybrid = methods.get("hybrid")
        fusion_text = ""
        if hybrid and "dense" in methods:
            rsig = sig_for("hybrid", "recall@100")
            rtext = ""
            if rsig:
                rtext = (f" (p = {rsig['p_value']:.4f}, significant)"
                         if rsig["significant_at_05"] else
                         f" (p = {rsig['p_value']:.4f})")
            fusion_text = (
                f" The clearest result is at the first stage: hybrid fusion lifts "
                f"Recall@100 to {hybrid['metrics']['recall@100']:.4f}, "
                f"{rsig['relative_pct']:+.1f}% over BM25{rtext} and above dense retrieval "
                f"({methods['dense']['metrics']['recall@100']:.4f}) too, so the two "
                f"retrievers fail on different queries. Since reranking cannot recover what "
                f"stage 1 missed, that recall is the ceiling stage 2 works within."
                if rsig else
                f" Hybrid fusion lifts Recall@100 to {hybrid['metrics']['recall@100']:.4f}, "
                f"above both BM25 ({base['metrics']['recall@100']:.4f}) and dense "
                f"({methods['dense']['metrics']['recall@100']:.4f}) alone.")

        if best_key == "bm25":
            headline = (
                f"Primary metric nDCG@{cut}: no configuration beat the BM25 baseline "
                f"({base['metrics'][f'ndcg@{cut}']:.4f}) on this query sample, which we "
                f"report as measured.")
        else:
            headline = (
                f"Primary metric nDCG@{cut}: the best configuration, {best['label']}, "
                f"reaches {best['metrics'][f'ndcg@{cut}']:.4f} against "
                f"{base['metrics'][f'ndcg@{cut}']:.4f} for BM25, a relative change of "
                f"{rel:+.1f}%. MRR moves from {base['metrics']['mrr']:.4f} to "
                f"{best['metrics']['mrr']:.4f} and MAP from {base['metrics']['map']:.4f} "
                f"to {best['metrics']['map']:.4f}.")
        body(doc, headline + sig_text + fusion_text)

        ce = methods.get("hybrid_ce_ft") or methods.get("hybrid_ce")
        cost_text = ""
        if ce:
            per_doc = ce["stage2_ms"] / max(ce["first_stage_k"], 1)
            full = per_doc * bench["corpus_size"] / 1000
            cost_text = (
                f" At {per_doc:.1f} ms per document, scoring all {bench['corpus_size']:,} for "
                f"one query would take ~{full:.0f} s; reranking the top "
                f"{ce['first_stage_k']} takes {ce['stage2_ms']/1000:.1f} s, "
                f"{full*1000/max(ce['stage2_ms'],1):.0f}x faster — the cascade's "
                f"quantitative justification.")

        body(doc,
             f"Evaluated on {bench['n_queries']} held-out test queries over "
             f"{bench['corpus_size']:,} documents." + cost_text +
             " Absolute scores are modest by nature: claim-to-prior-art pairs share very "
             "little vocabulary, and many relevant patents are never cited, so unlabelled "
             "correct answers count as errors. The comparison between methods is the "
             "meaningful result.")
    else:
        body(doc, "Run scripts/run_benchmark.py to populate this section.")

    doc.save(out_path)
    print(f"Wrote {out_path}")
    if bench:
        print(f"  best = {best['label']}  nDCG@{cut} = "
              f"{best['metrics'][f'ndcg@{cut}']:.4f}")


if __name__ == "__main__":
    build(ROOT / "Patent Prior Art Reranking Project Abstract.docx")
